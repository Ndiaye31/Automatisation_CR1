import openpyxl

# Charger le fichier
file_path = "planning.xlsx"
wb = openpyxl.load_workbook(file_path)
ws = wb["Feuil1"]

# Liste des valeurs à insérer
valeurs = [
    "Python", "Ckk", "Data", "Valeur 4",
    "Valeur 5", "Valeur 6", "Valeur 7", "Valeur 8",
    "Valeur 9", "Valeur 10", "Valeur 11", "Valeur 12",
    "Valeur 13"
]

# Lignes à remplir, en sautant 16 et 22
ligne = 12
val_index = 0
while val_index < len(valeurs):
    if ligne in (16, 22):  # sauter ces lignes
        ligne += 1
        continue

    cell = ws.cell(row=ligne, column=3)  # colonne C = 3
    # Vérifier si la cellule est fusionnée
    if cell.coordinate in ws.merged_cells:
        # Récupérer la cellule en haut à gauche du bloc fusionné
        for merged in ws.merged_cells.ranges:
            if cell.coordinate in merged:
                top_left = ws.cell(row=merged.min_row, column=merged.min_col)
                top_left.value = valeurs[val_index]
                break
    else:
        cell.value = valeurs[val_index]

    val_index += 1
    ligne += 1

# Sauvegarder le fichier
wb.save("Planning etudiant journalier mactar_rempli.xlsx")
